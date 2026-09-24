import os
import io
from datetime import datetime
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

# Optional dependency guards for PDF and Word export
HAS_REPORTLAB = False
HAS_DOCX = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        Image,
        KeepTogether,
        HRFlowable,
        PageBreak
    )
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    import docx
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def get_agronomic_advisory(features_df, crop_name):
    """
    Extract key soil and environmental parameters to provide field diagnosis and advisory.
    """
    feat_map = {}
    if isinstance(features_df, pd.DataFrame):
        if "Feature" in features_df.columns and "Representative Value" in features_df.columns:
            for _, r in features_df.iterrows():
                try:
                    feat_map[str(r["Feature"]).strip()] = float(str(r["Representative Value"]).strip())
                except ValueError:
                    feat_map[str(r["Feature"]).strip()] = str(r["Representative Value"]).strip()
        else:
            for col in features_df.columns:
                val = features_df[col].iloc[0]
                try:
                    feat_map[str(col).strip()] = float(val)
                except (ValueError, TypeError):
                    feat_map[str(col).strip()] = str(val)

    ph = feat_map.get("Soil pH Level", feat_map.get("Soil pH", 6.8))
    oc = feat_map.get("Organic Carbon (%)", 0.65)
    n_val = feat_map.get("Soil Nitrogen (N)", 12.0)
    p_val = feat_map.get("Soil Phosphorus (P)", 10.0)
    k_val = feat_map.get("Soil Potassium (K)", 15.0)
    rain = feat_map.get("Average Rainfall (mm)", 850.0)

    if ph < 6.0:
        ph_status = f"Acidic (pH {ph:.1f})"
        ph_note = "Lime or wood ash application recommended before sowing."
    elif ph <= 7.5:
        ph_status = f"Optimal (pH {ph:.1f})"
        ph_note = "Ideal neutral range for maximum nutrient bioavailability."
    else:
        ph_status = f"Alkaline (pH {ph:.1f})"
        ph_note = "Application of organic compost recommended to prevent micronutrient fixation."

    if oc >= 0.75:
        oc_status = f"High Fertility ({oc:.2f}%)"
    elif oc >= 0.50:
        oc_status = f"Medium Fertility ({oc:.2f}%)"
    else:
        oc_status = f"Low Fertility ({oc:.2f}%) - Compost needed"

    npk_status = f"N: {n_val:.1f} • P: {p_val:.1f} • K: {k_val:.1f}"

    if ph < 6.0:
        advisory = (
            f"For {crop_name.title()}, soil test values indicate mild acidity (pH {ph:.1f}). "
            "Incorporate agricultural lime during field preparation and follow split nitrogen dosing."
        )
    elif rain < 450:
        advisory = (
            f"For {crop_name.title()}, seasonal rainfall ({rain:.0f} mm) is relatively low. "
            "Ensure supplemental irrigation during the critical flowering and grain-filling stages."
        )
    elif oc < 0.50:
        advisory = (
            f"For {crop_name.title()}, organic carbon ({oc:.2f}%) indicates depleted organic matter. "
            "Apply well-decomposed farmyard manure or bio-fertilizers to improve water retention."
        )
    else:
        advisory = (
            f"Soil chemistry and moisture levels for {crop_name.title()} are well-balanced. "
            "Adhere to recommended regional fertilizer schedules and monitor moisture during reproductive growth."
        )

    return {
        "ph_status": ph_status,
        "ph_note": ph_note,
        "oc_status": oc_status,
        "npk_status": npk_status,
        "advisory": advisory
    }


def generate_charts_image_bytes(history, crop_yield_series, prediction, summary, crop, district):
    """
    Generate a high-resolution, publication-grade 4-panel statistical figure
    and return it as a BytesIO PNG stream for embedding into PDF and Word reports.
    """
    fig, axs = plt.subplots(2, 2, figsize=(10, 7.5), dpi=150, facecolor="#FFFFFF")
    fig.subplots_adjust(hspace=0.38, wspace=0.28, left=0.09, right=0.95, top=0.92, bottom=0.08)

    primary_color = "#1B4D2E"
    pred_val = prediction["predicted_yield"]

    # 1. Line Chart: Yield Over Time & Forecast
    ax1 = axs[0, 0]
    ax1.set_facecolor("#FAFCFA")
    years = history["Year"].values
    actuals = history["Yield"].values
    preds = history["Predicted"].values

    ax1.plot(years, actuals, marker="o", markersize=6, color="#2E7D32", linewidth=2.2, label="Actual Yield")
    ax1.plot(years, preds, marker="s", markersize=5, color="#1565C0", linestyle="--", linewidth=1.8, label="Model Fitted")

    next_yr = int(years[-1]) + 1
    ax1.plot([years[-1], next_yr], [actuals[-1], pred_val], color="#E65100", linestyle=":", linewidth=1.8)
    ax1.plot(next_yr, pred_val, marker="*", markersize=14, color="#E65100", label=f"Forecast ({next_yr})")

    ax1.annotate(
        f"Forecast\n{pred_val:.2f} t/ha",
        xy=(next_yr, pred_val),
        xytext=(0, 10),
        textcoords="offset points",
        ha="center",
        fontsize=7.5,
        fontweight="bold",
        color="#E65100",
        bbox=dict(boxstyle="round,pad=0.22", fc="#FFF3E0", ec="#E65100", lw=1)
    )

    ax1.set_xlim(years[0] - 0.4, next_yr + 0.6)
    ax1.set_title(f"Yield Over Time: {crop.title()} ({district.title()})", fontsize=10, fontweight="bold", color=primary_color)
    ax1.set_xlabel("Harvest Year", fontsize=8)
    ax1.set_ylabel("Yield (Tonnes / Ha)", fontsize=8)
    ax1.xaxis.set_major_locator(plt.MaxNLocator(integer=True))
    ax1.grid(True, linestyle=":", alpha=0.6)
    ax1.legend(fontsize=7, loc="best", framealpha=0.9)

    # 2. Histogram & Normal Curve: Yield Distribution
    ax2 = axs[0, 1]
    ax2.set_facecolor("#FAFCFA")
    sample_yields = crop_yield_series if len(crop_yield_series) > 10 else actuals
    n_bins = min(22, max(8, int(len(sample_yields) ** 0.5)))

    ax2.hist(
        sample_yields,
        bins=n_bins,
        color="#81C784",
        edgecolor="#2E7D32",
        alpha=0.65,
        density=True,
        label="Yield Frequency"
    )

    mu = float(np.mean(sample_yields))
    sigma = max(0.01, float(np.std(sample_yields)))
    x_axis = np.linspace(min(sample_yields), max(sample_yields), 100)
    pdf = (1.0 / (sigma * np.sqrt(2 * np.pi))) * np.exp(-0.5 * ((x_axis - mu) / sigma) ** 2)
    ax2.plot(x_axis, pdf, color="#1B5E20", linewidth=1.8, label="Normal Curve")

    ax2.axvline(pred_val, color="#D32F2F", linestyle="--", linewidth=2.0, label=f"Forecast: {pred_val:.2f} t/ha")
    ax2.axvline(mu, color="#1565C0", linestyle=":", linewidth=1.5, label=f"Mean: {mu:.2f} t/ha")

    diff_pct = ((pred_val - mu) / mu) * 100 if mu > 0 else 0
    ax2.text(
        0.04, 0.90,
        f"Benchmark: {'+' if diff_pct >= 0 else ''}{diff_pct:.1f}% vs Regional Mean",
        transform=ax2.transAxes,
        fontsize=7.5,
        fontweight="bold",
        color="#1B4D2E",
        bbox=dict(boxstyle="round,pad=0.25", fc="#E8F5E9", ec="#A5D6A7", lw=1)
    )

    ax2.set_title(f"Yield Distribution & Benchmark ({crop.title()})", fontsize=10, fontweight="bold", color=primary_color)
    ax2.set_xlabel("Yield (Tonnes / Ha)", fontsize=8)
    ax2.set_ylabel("Probability Density", fontsize=8)
    ax2.grid(True, linestyle=":", alpha=0.6)
    ax2.legend(fontsize=7, loc="upper right", framealpha=0.9)

    # 3. Bar Chart: Residual Accuracy Errors
    ax3 = axs[1, 0]
    ax3.set_facecolor("#FAFCFA")
    residuals = history["Residual"].values
    bar_colors = ["#2E7D32" if r >= 0 else "#C62828" for r in residuals]
    bars = ax3.bar(years, residuals, color=bar_colors, width=0.45, edgecolor="#333333", alpha=0.85)
    ax3.axhline(0, color="#1C251D", linestyle="-", linewidth=1.2)

    for bar, val in zip(bars, residuals):
        y_offset = 0.015 if val >= 0 else -0.035
        ax3.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + y_offset,
            f"{val:+.2f}",
            ha="center",
            va="bottom" if val >= 0 else "top",
            fontsize=7,
            fontweight="bold"
        )

    y_min, y_max = min(residuals), max(residuals)
    margin = max(0.12, max(abs(y_min), abs(y_max)) * 0.35)
    ax3.set_ylim(min(y_min - margin, -0.12), max(y_max + margin, 0.12))

    ax3.set_title("Historical Error Margins (Actual − Predicted)", fontsize=10, fontweight="bold", color=primary_color)
    ax3.set_xlabel("Harvest Year", fontsize=8)
    ax3.set_ylabel("Error Margin (t/ha)", fontsize=8)
    ax3.xaxis.set_major_locator(plt.MaxNLocator(integer=True))
    ax3.grid(True, linestyle=":", alpha=0.6)

    # 4. Donut Chart: Productivity Tiers
    ax4 = axs[1, 1]
    ax4.set_facecolor("#FAFCFA")
    tier_high = int((sample_yields >= 2.5).sum())
    tier_moderate = int(((sample_yields >= 1.5) & (sample_yields < 2.5)).sum())
    tier_low = int((sample_yields < 1.5).sum())

    tier_sizes = [tier_high, tier_moderate, tier_low]
    tier_labels = ["High (≥2.5 t/ha)", "Moderate (1.5–2.5)", "Low (<1.5 t/ha)"]
    tier_colors = ["#2E7D32", "#FFA000", "#D32F2F"]

    f_sizes, f_labels, f_colors = [], [], []
    for s, l, c in zip(tier_sizes, tier_labels, tier_colors):
        if s > 0:
            f_sizes.append(s)
            f_labels.append(l)
            f_colors.append(c)

    if f_sizes:
        wedges, texts, autotexts = ax4.pie(
            f_sizes,
            labels=f_labels,
            colors=f_colors,
            autopct="%1.1f%%",
            startangle=140,
            pctdistance=0.75,
            textprops={"fontsize": 7.5}
        )
        for autotext in autotexts:
            autotext.set_color("white")
            autotext.set_fontweight("bold")
            autotext.set_fontsize(7.5)

        centre_circle = plt.Circle((0, 0), 0.52, fc="white")
        ax4.add_artist(centre_circle)
        ax4.text(0, 0.05, f"{len(sample_yields)}", ha="center", va="center", fontsize=13, fontweight="bold", color=primary_color)
        ax4.text(0, -0.15, "Past Records", ha="center", va="center", fontsize=7, color="#5C6B5E")
        ax4.axis("equal")
    else:
        ax4.text(0.5, 0.5, "Data Unavailable", ha="center", va="center")

    ax4.set_title(f"Productivity Tier Breakdown ({crop.title()})", fontsize=10, fontweight="bold", color=primary_color)

    # Save to buffer
    img_buf = io.BytesIO()
    fig.savefig(img_buf, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    img_buf.seek(0)
    return img_buf


def export_pdf_report(filepath, crop, district, state, season, model_name, prediction, performance, summary, history, weather, features, crop_yield_series=None):
    """
    Generate an executive-grade PDF advisory report with embedded charts and tables.
    """
    if not HAS_REPORTLAB:
        raise ImportError(
            "PDF export requires the 'reportlab' package.\n"
            "Please install it using: pip install reportlab"
        )

    if crop_yield_series is None:
        crop_yield_series = np.array([])

    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        textColor=colors.HexColor("#1B4D2E"),
        spaceAfter=4
    )

    sub_style = ParagraphStyle(
        "ReportSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        textColor=colors.HexColor("#5C6B5E"),
        spaceAfter=12
    )

    section_style = ParagraphStyle(
        "SectionHead",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        textColor=colors.HexColor("#1B4D2E"),
        spaceBefore=10,
        spaceAfter=6
    )

    cell_bold = ParagraphStyle("CellBold", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, textColor=colors.HexColor("#1C251D"))
    cell_norm = ParagraphStyle("CellNorm", parent=styles["Normal"], fontName="Helvetica", fontSize=9, textColor=colors.HexColor("#333333"))
    cell_green = ParagraphStyle("CellGreen", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10, textColor=colors.HexColor("#2E7D32"))

    story = []

    # Title & Metadata Banner
    story.append(Paragraph(f"Crop Yield Report: {crop.title()}", title_style))
    meta_line = f"<b>Location:</b> {district.title()}, {state.title()} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Season:</b> {season.title()} Season &nbsp;&nbsp;|&nbsp;&nbsp; <b>Model:</b> {model_name} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Generated:</b> {datetime.now().strftime('%B %d, %Y')}"
    story.append(Paragraph(meta_line, sub_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#2E7D32"), spaceAfter=12))

    # Key Yield Prediction Table Card
    pred_val = prediction["predicted_yield"]
    quintals_ha = pred_val * 10.0
    kg_ha = pred_val * 1000.0

    if pred_val >= 3.0:
        tier_str = "High Yield"
        tier_color = colors.HexColor("#2E7D32")
    elif pred_val >= 1.8:
        tier_str = "Moderate Yield"
        tier_color = colors.HexColor("#F57F17")
    else:
        tier_str = "Low Yield"
        tier_color = colors.HexColor("#C62828")

    pred_data = [
        [
            Paragraph("<b>Predicted Yield</b>", cell_bold),
            Paragraph(f"<b>{pred_val:.3f} Tonnes / Ha</b><br/><font size=8 color='#5C6B5E'>({quintals_ha:.1f} Quintals/Ha • {kg_ha:.0f} kg/Ha)</font>", cell_green),
            Paragraph("<b>Expected Status</b>", cell_bold),
            Paragraph(f"<b>{tier_str}</b>", ParagraphStyle("Tier", parent=cell_bold, textColor=tier_color))
        ],
        [
            Paragraph("<b>Latest Recorded</b>", cell_bold),
            Paragraph(f"{prediction['actual_yield']:.3f} t/ha ({prediction['latest_year']})", cell_norm),
            Paragraph("<b>Prediction Error</b>", cell_bold),
            Paragraph(f"±{prediction['prediction_error']:.3f} t/ha", cell_norm)
        ],
        [
            Paragraph("<b>Model R² Score</b>", cell_bold),
            Paragraph(f"{performance['r2']:.4f}", cell_norm),
            Paragraph("<b>RMSE / MAE</b>", cell_bold),
            Paragraph(f"{performance['rmse']:.3f} / {performance['mae']:.3f} t/ha", cell_norm)
        ],
        [
            Paragraph("<b>Dataset Records</b>", cell_bold),
            Paragraph(f"{summary['crop_records']:,} records for this crop", cell_norm),
            Paragraph("<b>Sample Split</b>", cell_bold),
            Paragraph(f"{summary['training_samples']} train / {summary['testing_samples']} test", cell_norm)
        ]
    ]

    t_pred = Table(pred_data, colWidths=[1.4 * inch, 2.2 * inch, 1.4 * inch, 2.2 * inch])
    t_pred.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7FAF7")),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#D8E2D8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2ECE2")),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(t_pred)
    story.append(Spacer(1, 10))

    # Agronomic Field Diagnosis & Advisory
    diag = get_agronomic_advisory(features, crop)
    story.append(Paragraph("Agronomic Field Diagnosis & Advisory", section_style))
    diag_data = [
        [Paragraph("<b>Soil pH Diagnosis</b>", cell_bold), Paragraph(f"{diag['ph_status']} – {diag['ph_note']}", cell_norm)],
        [Paragraph("<b>Organic Carbon Status</b>", cell_bold), Paragraph(diag['oc_status'], cell_norm)],
        [Paragraph("<b>Primary Nutrients (N-P-K)</b>", cell_bold), Paragraph(diag['npk_status'], cell_norm)],
        [Paragraph("<b>Recommended Field Practice</b>", cell_bold), Paragraph(diag['advisory'], cell_norm)],
    ]
    t_diag = Table(diag_data, colWidths=[2.0 * inch, 5.2 * inch])
    t_diag.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7FAF7")),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#D8E2D8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2ECE2")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(t_diag)
    story.append(Spacer(1, 12))

    # Embedded Statistical Charts
    story.append(Paragraph("Statistical Visual Analytics", section_style))
    chart_buf = generate_charts_image_bytes(history, crop_yield_series, prediction, summary, crop, district)
    story.append(Image(chart_buf, width=7.2 * inch, height=5.4 * inch))
    story.append(Spacer(1, 10))
    story.append(PageBreak())

    # Historical Records Table (Page 2)
    story.append(Paragraph(f"Historical Yield Records: {district.title()} ({crop.title()})", section_style))

    hist_headers = [Paragraph(f"<b>{col}</b>", ParagraphStyle("TH", parent=cell_bold, textColor=colors.white)) for col in history.columns]
    hist_rows = [hist_headers]

    for idx, (_, row) in enumerate(history.iterrows()):
        row_cells = []
        for val in row:
            if isinstance(val, float):
                row_cells.append(Paragraph(f"{val:.3f}", cell_norm))
            else:
                row_cells.append(Paragraph(str(val), cell_norm))
        hist_rows.append(row_cells)

    col_w = (7.2 * inch) / len(history.columns)
    t_hist = Table(hist_rows, colWidths=[col_w] * len(history.columns))
    t_hist.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B4D2E")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F9FAF9")]),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#D8E2D8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5EBE5")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t_hist)
    story.append(Spacer(1, 14))

    # Historical Climate & Soil Readings (if available)
    if weather is not None and isinstance(weather, pd.DataFrame) and not weather.empty:
        story.append(Paragraph("Historical Climate & Soil Readings (Recent Seasons)", section_style))
        w_headers = [Paragraph(f"<b>{col}</b>", ParagraphStyle("THW", parent=cell_bold, fontSize=7.5, textColor=colors.white)) for col in weather.columns]
        w_rows = [w_headers]
        for _, row in weather.head(8).iterrows():
            w_cells = []
            for val in row:
                if isinstance(val, float):
                    w_cells.append(Paragraph(f"{val:.2f}" if abs(val) < 100 else f"{val:.0f}", ParagraphStyle("CW", parent=cell_norm, fontSize=7.5)))
                else:
                    w_cells.append(Paragraph(str(val), ParagraphStyle("CW", parent=cell_norm, fontSize=7.5)))
            w_rows.append(w_cells)
        col_w_w = (7.2 * inch) / len(weather.columns)
        t_weather = Table(w_rows, colWidths=[col_w_w] * len(weather.columns))
        t_weather.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B4D2E")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F9FAF9")]),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#D8E2D8")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5EBE5")),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(t_weather)
        story.append(Spacer(1, 14))

    # Soil Features & Field Parameters Profile
    story.append(Paragraph("Soil Features & Field Parameters Profile", section_style))
    feat_headers = [Paragraph("<b>Parameter</b>", ParagraphStyle("TH", parent=cell_bold, textColor=colors.white)),
                    Paragraph("<b>Value Used</b>", ParagraphStyle("TH", parent=cell_bold, textColor=colors.white))]
    feat_rows = [feat_headers]

    if "Feature" in features.columns and "Representative Value" in features.columns:
        for _, row in features.iterrows():
            feat_rows.append([Paragraph(str(row["Feature"]), cell_norm), Paragraph(str(row["Representative Value"]), cell_bold)])
    else:
        for col in features.columns:
            val = features[col].iloc[0]
            val_str = f"{val:.3f}" if isinstance(val, float) else str(val)
            feat_rows.append([Paragraph(str(col), cell_norm), Paragraph(val_str, cell_bold)])

    t_feat = Table(feat_rows, colWidths=[3.8 * inch, 3.4 * inch])
    t_feat.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D32")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F9FAF9")]),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#D8E2D8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5EBE5")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t_feat)

    # Build PDF
    doc.build(story)
    return filepath


def export_docx_report(filepath, crop, district, state, season, model_name, prediction, performance, summary, history, weather, features, crop_yield_series=None):
    """
    Generate a rich Microsoft Word (.docx) advisory report with embedded charts and tables.
    """
    if not HAS_DOCX:
        raise ImportError(
            "Word document export requires the 'python-docx' package.\n"
            "Please install it using: pip install python-docx"
        )

    if crop_yield_series is None:
        crop_yield_series = np.array([])

    doc = Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # Document Title
    title = doc.add_heading(f"Crop Yield Report: {crop.title()}", level=0)
    for run in title.runs:
        run.font.name = "Segoe UI"
        run.font.color.rgb = RGBColor(0x1B, 0x4D, 0x2E)

    # Subtitle Metadata
    sub = doc.add_paragraph()
    sub_run = sub.add_run(
        f"Location: {district.title()}, {state.title()}  •  Season: {season.title()} Season  •  Model: {model_name}\n"
        f"Generated: {datetime.now().strftime('%B %d, %Y - %I:%M %p')}"
    )
    sub_run.font.name = "Segoe UI"
    sub_run.font.size = Pt(9.5)
    sub_run.font.color.rgb = RGBColor(0x5C, 0x6B, 0x5E)

    # 1. Executive Summary Table
    doc.add_heading("Executive Summary & Prediction Details", level=1)
    pred_val = prediction["predicted_yield"]
    quintals_ha = pred_val * 10.0
    kg_ha = pred_val * 1000.0

    if pred_val >= 3.0:
        tier_str = "High Yield"
    elif pred_val >= 1.8:
        tier_str = "Moderate Yield"
    else:
        tier_str = "Low Yield"

    table_data = [
        ("Predicted Crop Yield", f"{pred_val:.3f} Tonnes / Hectare ({quintals_ha:.1f} Quintals/Ha, {kg_ha:.0f} kg/Ha)"),
        ("Yield Potential Status", tier_str),
        ("Latest Recorded Yield", f"{prediction['actual_yield']:.3f} t/ha (Harvest Year: {prediction['latest_year']})"),
        ("Validation Prediction Error", f"±{prediction['prediction_error']:.3f} t/ha"),
        ("Model Goodness of Fit (R²)", f"{performance['r2']:.4f}"),
        ("Root Mean Squared Error (RMSE)", f"{performance['rmse']:.4f} t/ha"),
        ("Mean Absolute Error (MAE)", f"{performance['mae']:.4f} t/ha"),
        ("Total Records for Crop", f"{summary['crop_records']:,} verified observations"),
        ("Sample Partition", f"{summary['training_samples']} training samples / {summary['testing_samples']} test samples")
    ]

    t = doc.add_table(rows=len(table_data) + 1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"

    hdr_cells = t.rows[0].cells
    hdr_cells[0].text = "Metric / Parameter"
    hdr_cells[1].text = "Value"

    for i in range(2):
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.name = "Segoe UI"
                run.font.color.rgb = RGBColor(0x1B, 0x4D, 0x2E)

    for idx, (label, val) in enumerate(table_data):
        row_cells = t.rows[idx + 1].cells
        row_cells[0].text = label
        row_cells[1].text = str(val)
        for cell in row_cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Segoe UI"
                    r.font.size = Pt(9.5)

    doc.add_paragraph()

    # Agronomic Field Diagnosis & Advisory
    diag = get_agronomic_advisory(features, crop)
    doc.add_heading("Agronomic Field Diagnosis & Advisory", level=1)
    diag_rows = [
        ("Soil pH Diagnosis", f"{diag['ph_status']} – {diag['ph_note']}"),
        ("Organic Carbon Status", diag['oc_status']),
        ("Primary Nutrients (N-P-K)", diag['npk_status']),
        ("Recommended Field Practice", diag['advisory'])
    ]
    t_diag = doc.add_table(rows=len(diag_rows) + 1, cols=2)
    t_diag.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_diag.style = "Table Grid"
    t_diag.rows[0].cells[0].text = "Field Indicator"
    t_diag.rows[0].cells[1].text = "Diagnosis / Recommendation"
    for cell in t_diag.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.name = "Segoe UI"
                r.font.color.rgb = RGBColor(0x1B, 0x4D, 0x2E)

    for idx, (label, val) in enumerate(diag_rows):
        cells = t_diag.rows[idx + 1].cells
        cells[0].text = label
        cells[1].text = val
        for cell in cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Segoe UI"
                    r.font.size = Pt(9.5)

    doc.add_paragraph()

    # 2. Visual Charts Section
    doc.add_heading("Visual Analytics & Statistical Charts", level=1)
    chart_buf = generate_charts_image_bytes(history, crop_yield_series, prediction, summary, crop, district)
    doc.add_picture(chart_buf, width=Inches(6.6))
    doc.add_page_break()

    # 3. Historical Cultivation Records
    doc.add_heading("Historical Harvest Records", level=1)
    hist_cols = list(history.columns)
    t_hist = doc.add_table(rows=len(history) + 1, cols=len(hist_cols))
    t_hist.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hist.style = "Table Grid"

    for i, col in enumerate(hist_cols):
        t_hist.rows[0].cells[i].text = str(col)
        for p in t_hist.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.name = "Segoe UI"

    for row_idx, (_, row) in enumerate(history.iterrows()):
        cells = t_hist.rows[row_idx + 1].cells
        for col_idx, val in enumerate(row):
            formatted_val = f"{val:.3f}" if isinstance(val, float) else str(val)
            cells[col_idx].text = formatted_val
            for p in cells[col_idx].paragraphs:
                for r in p.runs:
                    r.font.name = "Segoe UI"
                    r.font.size = Pt(9)

    doc.add_paragraph()

    # 4. Climate & Soil Readings (if available)
    if weather is not None and isinstance(weather, pd.DataFrame) and not weather.empty:
        doc.add_paragraph()
        doc.add_heading("Historical Climate & Soil Readings (Recent Seasons)", level=1)
        w_cols = list(weather.columns)
        t_w = doc.add_table(rows=min(9, len(weather) + 1), cols=len(w_cols))
        t_w.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_w.style = "Table Grid"

        for i, col in enumerate(w_cols):
            t_w.rows[0].cells[i].text = str(col)
            for p in t_w.rows[0].cells[i].paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.name = "Segoe UI"
                    r.font.size = Pt(8.5)

        for row_idx, (_, row) in enumerate(weather.head(8).iterrows()):
            cells = t_w.rows[row_idx + 1].cells
            for col_idx, val in enumerate(row):
                formatted_val = f"{val:.2f}" if isinstance(val, float) else str(val)
                cells[col_idx].text = formatted_val
                for p in cells[col_idx].paragraphs:
                    for r in p.runs:
                        r.font.name = "Segoe UI"
                        r.font.size = Pt(8.5)

    doc.add_paragraph()

    # 5. Soil Features & Field Parameters Profile
    doc.add_heading("Soil Features & Field Parameters Profile", level=1)
    if "Feature" in features.columns and "Representative Value" in features.columns:
        t_feat = doc.add_table(rows=len(features) + 1, cols=2)
        t_feat.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_feat.style = "Table Grid"
        t_feat.rows[0].cells[0].text = "Parameter"
        t_feat.rows[0].cells[1].text = "Value"
        for cell in t_feat.rows[0].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.name = "Segoe UI"

        for idx, (_, row) in enumerate(features.iterrows()):
            cells = t_feat.rows[idx + 1].cells
            cells[0].text = str(row["Feature"])
            cells[1].text = str(row["Representative Value"])
            for cell in cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = "Segoe UI"
                        r.font.size = Pt(9)
    else:
        feat_cols = list(features.columns)
        t_feat = doc.add_table(rows=len(feat_cols) + 1, cols=2)
        t_feat.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_feat.style = "Table Grid"
        t_feat.rows[0].cells[0].text = "Parameter"
        t_feat.rows[0].cells[1].text = "Value"
        for cell in t_feat.rows[0].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.name = "Segoe UI"

        for idx, col in enumerate(feat_cols):
            val = features[col].iloc[0]
            val_str = f"{val:.3f}" if isinstance(val, float) else str(val)
            cells = t_feat.rows[idx + 1].cells
            cells[0].text = str(col)
            cells[1].text = val_str
            for cell in cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = "Segoe UI"
                        r.font.size = Pt(9)

    doc.save(filepath)
    return filepath


def export_txt_report(filepath, crop, district, state, season, model_name, prediction, performance, summary, history, weather, features):
    """
    Generate clean plaintext summary report as fallback.
    """
    diag = get_agronomic_advisory(features, crop)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write("=" * 65 + "\n")
        f.write("CROP YIELD PREDICTION REPORT\n")
        f.write("=" * 65 + "\n\n")
        f.write(f"Crop:               {crop.title()}\n")
        f.write(f"Location:           {district.title()}, {state.title()}\n")
        f.write(f"Season:             {season.title()} Season\n")
        f.write(f"Model:              {model_name}\n")
        f.write(f"Predicted Yield:    {prediction['predicted_yield']:.3f} Tonnes / Hectare\n")
        f.write(f"Latest Recorded:    {prediction['actual_yield']:.3f} Tonnes / Hectare ({prediction['latest_year']})\n")
        f.write(f"Prediction Error:   {prediction['prediction_error']:.3f} Tonnes / Hectare\n\n")
        f.write("-" * 65 + "\n")
        f.write("AGRONOMIC FIELD DIAGNOSIS & ADVISORY\n")
        f.write("-" * 65 + "\n")
        f.write(f"Soil pH Status:     {diag['ph_status']} ({diag['ph_note']})\n")
        f.write(f"Organic Carbon:     {diag['oc_status']}\n")
        f.write(f"Nutrient Balance:   {diag['npk_status']}\n")
        f.write(f"Field Advisory:     {diag['advisory']}\n\n")
        f.write("-" * 65 + "\n")
        f.write("MODEL PERFORMANCE\n")
        f.write("-" * 65 + "\n")
        f.write(f"R² Score:           {performance['r2']:.4f}\n")
        f.write(f"RMSE:               {performance['rmse']:.4f}\n")
        f.write(f"MAE:                {performance['mae']:.4f}\n")
        f.write(f"Training Samples:   {summary['training_samples']}\n")
        f.write(f"Testing Samples:    {summary['testing_samples']}\n\n")
        f.write("-" * 65 + "\n")
        f.write("HISTORICAL HARVEST RECORDS\n")
        f.write("-" * 65 + "\n")
        f.write(history.to_string(index=False))
        if weather is not None and isinstance(weather, pd.DataFrame) and not weather.empty:
            f.write("\n\n" + "-" * 65 + "\n")
            f.write("HISTORICAL CLIMATE & SOIL READINGS\n")
            f.write("-" * 65 + "\n")
            f.write(weather.to_string(index=False))
        f.write("\n\n" + "-" * 65 + "\n")
        f.write("SOIL FEATURES & FIELD PARAMETERS PROFILE\n")
        f.write("-" * 65 + "\n")
        f.write(features.to_string(index=False))
        f.write("\n\n" + "=" * 65 + "\n")
    return filepath
